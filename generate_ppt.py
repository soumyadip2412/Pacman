from pptx import Presentation
from pptx.util import Inches, Pt

slides = [
    ("Intelligent Pac-Man Agent Using A* Search Algorithm", "Student Name \nGuide: [Guide Name] \nDepartment of CSE"),
    ("Motivation", "Importance of AI in games\nProject objectives: A* implementation, evaluation, visualization"),
    ("Problem Statement", "Collect all pellets while avoiding ghosts. Real-time planning and replanning in grid maze."),
    ("Agent Design", "Goal-based agent\nPerceive -> Plan -> Act cycle\nState: (row, col)") ,
    ("Algorithms", "BFS, DFS, and A* (f(n) = g(n) + h(n))\nHeuristic: Manhattan distance"),
    ("Heuristic", "Manhattan distance: h(n)=|r1-r2|+|c1-c2|\nAdmissible and consistent for 4-dir grids"),
    ("System Architecture", "Modules: main.py, pacman.py, astar.py, bfs.py, dfs.py, maze.py\nRenderer: Pygame; Metrics: AlgorithmMetrics"),
    ("Results", "Nodes explored (A* < BFS)<br/>Path cost: A* optimal<br/>Execution time: A* efficient"),
    ("Discussion", "Why A* performed best\nTrade-offs: replanning frequency vs search cost"),
    ("Conclusion & Future Work", "Summary of achievements\nFuture: RL, adaptive ghosts, dynamic mazes")
]

prs = Presentation()
# Use title slide layout for slide 1
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = slides[0][0]
subtitle.text = slides[0][1]

for title_text, body_text in slides[1:]:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = title_text
    tf = body.text_frame
    # split body_text into lines by newline or <br/>
    lines = body_text.replace('<br/>', '\n').split('\n')
    tf.text = lines[0]
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 1

out_path = 'slides.pptx'
prs.save(out_path)
print(f"Generated {out_path}")
